from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt


def build_pptx(title: str, outputs: dict) -> bytes:
    presentation = outputs.get("presentation", {})
    slides = presentation.get("slides", [])
    if not slides:
        # Make a valid presentation even if the user exports without a presentation output.
        slides = [{"title": title, "bullets": ["No presentation output was generated."]}]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for item in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        title_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(0.55), Inches(12), Inches(0.8)
        )
        p = title_box.text_frame.paragraphs[0]
        p.text = str(item.get("title", "Slide"))
        p.font.size = Pt(28)
        p.font.bold = True

        body = slide.shapes.add_textbox(
            Inches(0.9), Inches(1.6), Inches(11.5), Inches(5.2)
        )
        tf = body.text_frame
        tf.clear()
        bullets = item.get("bullets", [])
        for i, bullet in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = str(bullet)
            para.font.size = Pt(20)
            para.level = 0
            para.space_after = Pt(10)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
